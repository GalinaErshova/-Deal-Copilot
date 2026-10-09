"""Document processing layer adapted from the Ask-Learn ingestion approach.

The important contract is stable: every parser returns ordered blocks with a
source path/page so the UI can compare ORIGINAL vs PARSED and later trace AI
fields back to evidence.
"""
from __future__ import annotations

import email
import email.policy
import json
import re
import zipfile
from dataclasses import asdict, dataclass
from html.parser import HTMLParser
from io import BytesIO
from itertools import islice
from pathlib import Path
from typing import ClassVar


@dataclass
class ParsedBlock:
    kind: str
    text: str
    path: str
    page_no: int | None = None
    title: str | None = None
    rows: list[list[str]] | None = None

@dataclass
class ParsedDocument:
    parser: str
    confidence: str
    blocks: list[ParsedBlock]
    warnings: list[str]

    def to_json(self) -> str:
        return json.dumps(
            {"parser":self.parser,"confidence":self.confidence,
             "blocks":[asdict(x) for x in self.blocks],"warnings":self.warnings},
            ensure_ascii=False,
        )

    @property
    def plain_text(self) -> str:
        return "\n\n".join(x.text for x in self.blocks if x.text)

    def extraction_text(self, filename: str) -> str:
        """Text for LLM with stable evidence markers.

        MiMo must copy source_document/source_location from these markers,
        so every extracted field can be traced back to a parsed block.
        """
        parts: list[str] = []
        for block in self.blocks:
            if not block.text:
                continue
            page = f" page={block.page_no}" if block.page_no else ""
            parts.append(
                f"[DOCUMENT={filename} PATH={block.path}{page} KIND={block.kind}]\n{block.text}"
            )
        return "\n\n".join(parts)


def chunk_extraction_text(text: str, max_chars: int) -> list[str]:
    """Разбивает текст для небольшой локальной модели, сохраняя маркеры источников."""
    if max_chars <= 0:
        raise ValueError("max_chars должен быть больше нуля")
    if len(text) <= max_chars:
        return [text] if text else []

    sections = re.split(r"(?=\[DOCUMENT=)", text)
    pieces: list[str] = []
    for section in sections:
        if not section:
            continue
        header, separator, body = section.partition("\n")
        prefix = f"{header}{separator}" if separator else ""
        body_limit = max_chars - len(prefix)
        if body_limit <= 0:
            raise ValueError("max_chars слишком мало для маркера документа")

        while body:
            if len(body) <= body_limit:
                fragment, body = body, ""
            else:
                boundary = max(body.rfind("\n", 0, body_limit), body.rfind(" ", 0, body_limit))
                if boundary < body_limit // 2:
                    boundary = body_limit
                fragment, body = body[:boundary].rstrip(), body[boundary:].lstrip()
            if fragment:
                pieces.append(prefix + fragment)

    chunks: list[str] = []
    current = ""
    for piece in pieces:
        candidate = f"{current}\n\n{piece}" if current else piece
        if current and len(candidate) > max_chars:
            chunks.append(current)
            current = piece
        else:
            current = candidate
    if current:
        chunks.append(current)
    return chunks

def _cell(value: object) -> str:
    return "" if value is None else str(value)


def _table_text(rows: list[list[str]]) -> str:
    """Делает строки таблицы самодостаточными: заголовок столбца → значение."""
    rows = [[" ".join(cell.split()) for cell in row] for row in rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        return ""

    header = rows[0] if len(rows) > 1 else []
    data_rows = rows[1:] if header else rows
    lines = []
    for row_number, row in enumerate(data_rows, start=2 if header else 1):
        cells = []
        for column, value in enumerate(row):
            if not value:
                continue
            title = header[column].strip() if column < len(header) else ""
            cells.append(f"{title or f'Столбец {column + 1}'}: {value}")
        if cells:
            lines.append(f"Строка {row_number}: " + "; ".join(cells))
    return "\n".join(lines)


def _paragraph_blocks(text: str, kind: str = "paragraph", prefix: str = "/p") -> list[ParsedBlock]:
    """Разбивает сплошной текст на абзацы с устойчивыми путями-источниками."""
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    return [
        ParsedBlock(kind, value, f"{prefix}/{index}")
        for index, value in enumerate((item.strip() for item in re.split(r"\n\s*\n", normalized)), start=1)
        if value
    ]


def _decode_text(data: bytes) -> tuple[str, str]:
    """Выбирает кодировку текста, не превращая русские файлы Windows-1251 в кракозябры."""
    for encoding in ("utf-8-sig", "utf-8", "cp1251", "koi8-r"):
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), "utf-8 (с заменами)"


class _HtmlCollector(HTMLParser):
    """Извлекает видимый текст и таблицы HTML без выполнения содержимого."""

    hidden_tags: ClassVar[frozenset[str]] = frozenset({"script", "style", "title", "noscript", "template"})

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks: list[ParsedBlock] = []
        self.text: list[str] = []
        self.current_tag: str | None = None
        self.hidden_depth = 0
        self.table_rows: list[list[str]] | None = None
        self.row: list[str] | None = None
        self.cell: list[str] | None = None

    def _flush_text(self):
        value = " ".join(" ".join(self.text).split())
        if value:
            index = len(self.blocks) + 1
            kind = "heading" if self.current_tag and self.current_tag.startswith("h") else "paragraph"
            self.blocks.append(ParsedBlock(kind, value, f"/n/{index}", title=value if kind == "heading" else None))
        self.text = []
        self.current_tag = None

    def handle_starttag(self, tag, attrs):
        if tag in self.hidden_tags:
            self.hidden_depth += 1
        elif not self.hidden_depth and tag == "table":
            self._flush_text()
            self.table_rows = []
        elif not self.hidden_depth and tag == "tr" and self.table_rows is not None:
            self.row = []
        elif not self.hidden_depth and tag in {"td", "th"} and self.row is not None:
            self.cell = []
        elif not self.hidden_depth and tag in {"p", "li", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self._flush_text()
            self.current_tag = tag

    def handle_endtag(self, tag):
        if tag in self.hidden_tags:
            self.hidden_depth = max(0, self.hidden_depth - 1)
        elif tag in {"td", "th"} and self.cell is not None and self.row is not None:
            self.row.append(" ".join(" ".join(self.cell).split()))
            self.cell = None
        elif tag == "tr" and self.row is not None and self.table_rows is not None:
            self.table_rows.append(self.row)
            self.row = None
        elif tag == "table" and self.table_rows is not None:
            rows = self.table_rows
            self.table_rows = None
            text = _table_text(rows)
            if text:
                self.blocks.append(ParsedBlock("table", text, f"/table/{len(self.blocks) + 1}", rows=rows))
        elif tag in {"p", "li", "h1", "h2", "h3", "h4", "h5", "h6", "div", "br"} and not self.hidden_depth:
            self._flush_text()

    def handle_data(self, data):
        if self.hidden_depth:
            return
        if self.cell is not None:
            self.cell.append(data)
        elif self.table_rows is None:
            self.text.append(data)


def _parse_html(text: str) -> ParsedDocument:
    collector = _HtmlCollector()
    collector.feed(text)
    collector._flush_text()
    return ParsedDocument("html", "high" if collector.blocks else "low", collector.blocks, [])


def _parse_markdown(text: str) -> ParsedDocument:
    """Сохраняет заголовки, абзацы и Markdown-таблицы отдельными блоками."""
    lines = text.splitlines()
    blocks: list[ParsedBlock] = []
    paragraph: list[str] = []
    index = 0

    def flush():
        value = " ".join(" ".join(paragraph).split())
        if value:
            blocks.append(ParsedBlock("paragraph", value, f"/p/{len(blocks) + 1}"))
        paragraph.clear()

    while index < len(lines):
        line = lines[index]
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            flush()
            title = heading.group(2).strip()
            blocks.append(ParsedBlock("heading", title, f"/h/{len(blocks) + 1}", title=title))
        elif "|" in line and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{3,}", lines[index + 1]):
            flush()
            rows = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in [line]]
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            rows = [row for row in rows if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in row)]
            value = _table_text(rows)
            if value:
                blocks.append(ParsedBlock("table", value, f"/table/{len(blocks) + 1}", rows=rows))
            continue
        elif line.strip():
            paragraph.append(line.strip())
        else:
            flush()
        index += 1
    flush()
    return ParsedDocument("markdown", "high" if blocks else "low", blocks, [])

def parse_xlsx(data: bytes, max_rows: int) -> ParsedDocument:
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    blocks: list[ParsedBlock] = []
    warnings: list[str] = []
    for index, sheet in enumerate(book.worksheets, start=1):
        raw_rows = list(islice(sheet.iter_rows(values_only=True), max_rows + 1))
        truncated = len(raw_rows) > max_rows
        rows = [[_cell(v) for v in row] for row in raw_rows[:max_rows]]
        rows = [r for r in rows if any(c.strip() for c in r)]
        if not rows:
            continue
        md = _table_text(rows)
        blocks.append(ParsedBlock("sheet", md, f"/sheet/{index}", title=sheet.title, rows=rows))
        if len(rows) > 1:
            warnings.append(f"Лист {sheet.title}: первая строка принята за заголовок столбцов — сверьте с оригиналом")
        if truncated:
            warnings.append(f"Лист {sheet.title}: обработаны первые {max_rows} строк")
    return ParsedDocument("xlsx","medium",blocks,warnings)

def parse_docx(data: bytes) -> ParsedDocument:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    document = Document(BytesIO(data))
    blocks: list[ParsedBlock] = []
    warnings: list[str] = []
    body = document.element.body
    for order, element in enumerate(body.iterchildren(), start=1):
        tag = element.tag.rsplit("}",1)[-1]
        if tag == "tbl":
            table = Table(element, document)
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            text = _table_text(rows)
            blocks.append(ParsedBlock("table",text,f"/table/{order}",rows=rows))
            if len(rows) > 1:
                warnings.append(f"Таблица {order}: первая строка принята за заголовок столбцов — сверьте с оригиналом")
        elif tag == "p":
            p = Paragraph(element, document)
            text = " ".join(p.text.split())
            if text:
                style = (p.style.name if p.style else "").lower()
                blocks.append(ParsedBlock("heading" if style.startswith("heading") else "paragraph",
                                          text,f"/p/{order}",title=text if style.startswith("heading") else None))
    return ParsedDocument("docx","high" if blocks else "low",blocks,warnings)

def parse_pdf(data: bytes) -> ParsedDocument:
    import fitz
    doc = fitz.open(stream=data, filetype="pdf")
    blocks: list[ParsedBlock] = []
    warnings: list[str] = []
    scanned_pages = 0
    for page_index, page in enumerate(doc, start=1):
        table_frames = []
        try:
            table_frames = list(page.find_tables().tables)
        except Exception:  # noqa: BLE001 — распознавание геометрии таблиц эвристическое
            warnings.append(f"Стр. {page_index}: не удалось проверить структуру таблиц")
        table_bounds = []
        positioned: list[tuple[float, ParsedBlock]] = []
        for table_index, frame in enumerate(table_frames, start=1):
            rows = [[_cell(value) for value in row] for row in frame.extract()]
            rows = [[cell.strip() for cell in row] for row in rows]
            text = _table_text(rows)
            if text:
                table_bounds.append(tuple(frame.bbox))
                positioned.append((frame.bbox[1], ParsedBlock(
                    "table", text, f"/page/{page_index}/t/{table_index}", page_no=page_index, rows=rows
                )))
                warnings.append(f"Стр. {page_index}, таблица {table_index}: проверьте заголовок и объединённые ячейки")
        text_blocks = [block for block in page.get_text("blocks") if block[6] == 0]
        page_text = []
        for idx, block in enumerate(text_blocks, start=1):
            text = " ".join(str(block[4]).split())
            bounds = tuple(block[:4])
            overlaps_table = any(
                bounds[0] >= frame[0] and bounds[1] >= frame[1]
                and bounds[2] <= frame[2] and bounds[3] <= frame[3]
                for frame in table_bounds
            )
            if text and not overlaps_table:
                page_text.append(text)
                positioned.append((block[1], ParsedBlock("paragraph",text,f"/page/{page_index}/p/{idx}",page_no=page_index)))
        if not positioned:
            scanned_pages += 1
            blocks.append(ParsedBlock("scan","",f"/page/{page_index}/scan",page_no=page_index))
        else:
            blocks.extend(block for _top, block in sorted(positioned, key=lambda item: item[0]))
    if scanned_pages:
        warnings.append(f"{scanned_pages} стр. без текстового слоя: требуется OCR")
    confidence = "high" if blocks and not scanned_pages else "medium" if blocks else "low"
    return ParsedDocument("pdf",confidence,blocks,warnings)


def parse_pptx(data: bytes) -> ParsedDocument:
    """Извлекает текст, таблицы и номера слайдов презентации PPTX."""
    from pptx import Presentation

    presentation = Presentation(BytesIO(data))
    blocks: list[ParsedBlock] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        for shape_number, shape in enumerate(slide.shapes, start=1):
            path = f"/slide/{slide_number}/shape/{shape_number}"
            if getattr(shape, "has_table", False):
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                text = _table_text(rows)
                if text:
                    blocks.append(ParsedBlock("table", text, path, page_no=slide_number, rows=rows))
            elif getattr(shape, "has_text_frame", False):
                text = "\n".join(" ".join(p.text.split()) for p in shape.text_frame.paragraphs).strip()
                if text:
                    blocks.append(ParsedBlock("slide", text, path, page_no=slide_number))
    return ParsedDocument("pptx", "high" if blocks else "low", blocks, [])


_RTF_TOKEN = re.compile(rb"\\([a-zA-Z]+)(-?\d+)? ?|\\'([0-9a-fA-F]{2})|\\(.)|([{}])|[\r\n]+|([^\\{}\r\n]+)")
_RTF_SKIP = {"fonttbl", "colortbl", "stylesheet", "info", "pict", "object", "header", "footer", "listtable", "listoverridetable"}
_RTF_WORDS = {"par":"\n\n", "line":"\n", "tab":"\t", "cell":"\t", "row":"\n", "emdash":"—", "endash":"–", "bullet":"•"}


def _rtf_text(data: bytes) -> str:
    """Удаляет управляющую разметку RTF и сохраняет Unicode/Windows-кодировку."""
    output: list[str] = []
    pending = bytearray()
    codec = "cp1251"
    skip = False
    uc = 1
    to_skip = 0
    group_start = False
    stack: list[tuple[bool, int]] = []

    def flush():
        if pending:
            output.append(pending.decode(codec, errors="replace"))
            pending.clear()

    for match in _RTF_TOKEN.finditer(data):
        word, argument, hex_byte, symbol, brace, text = match.groups()
        starts_group = False
        if brace == b"{":
            flush(); stack.append((skip, uc)); starts_group = True; to_skip = 0
        elif brace == b"}":
            flush()
            if stack:
                skip, uc = stack.pop()
            to_skip = 0
        elif word is not None:
            name = word.decode("ascii")
            if group_start and name in _RTF_SKIP:
                skip = True
            if name == "ansicpg" and argument:
                flush()
                try:
                    codec = f"cp{int(argument)}"
                    "".encode(codec)
                except (ValueError, LookupError):
                    codec = "cp1251"
            elif name == "uc" and argument:
                uc = int(argument)
            elif name == "u" and argument:
                flush()
                if not skip:
                    codepoint = int(argument)
                    output.append(chr(codepoint + 65536 if codepoint < 0 else codepoint))
                to_skip = uc
            elif name in _RTF_WORDS and not skip:
                flush(); output.append(_RTF_WORDS[name])
        elif hex_byte is not None:
            if to_skip:
                to_skip -= 1
            elif not skip:
                pending.append(int(hex_byte, 16))
        elif symbol is not None:
            if symbol == b"*":
                skip = skip or group_start; starts_group = group_start
            elif to_skip:
                to_skip -= 1
            elif not skip:
                flush()
                if symbol in (b"\\", b"{", b"}"):
                    output.append(symbol.decode("ascii"))
                elif symbol == b"~":
                    output.append("\u00a0")
                elif symbol == b"_":
                    output.append("-")
        elif text is not None:
            chunk = text
            if to_skip:
                dropped = min(to_skip, len(chunk)); chunk = chunk[dropped:]; to_skip -= dropped
            if chunk and not skip:
                flush(); output.append(chunk.decode(codec, errors="replace"))
        group_start = starts_group
    flush()
    return "".join(output)


def parse_rtf(data: bytes) -> ParsedDocument:
    if not data.lstrip().startswith(b"{\\rtf"):
        raise ValueError("Файл не содержит сигнатуру RTF")
    blocks = _paragraph_blocks(_rtf_text(data))
    return ParsedDocument("rtf", "medium", blocks, ["RTF преобразован в текст; исходное форматирование могло быть упрощено"])


def parse_mht(data: bytes) -> ParsedDocument:
    """Достаёт HTML-часть MHT/MHTML и передаёт её безопасному HTML-парсеру."""
    message = email.message_from_bytes(data, policy=email.policy.default)
    for part in message.walk():
        if part.get_content_type() != "text/html":
            continue
        payload = part.get_payload(decode=True)
        if not isinstance(payload, bytes):
            continue
        charset = part.get_content_charset() or "cp1251"
        try:
            html = payload.decode(charset, errors="replace")
        except LookupError:
            html = payload.decode("cp1251", errors="replace")
        parsed = _parse_html(html)
        parsed.parser = "mht"
        parsed.warnings.append("Извлечён текст HTML-части MHT; вложенные изображения не распознавались")
        return parsed
    return ParsedDocument("mht", "low", [], ["В MHT-файле не найдена HTML-часть"])

def parse_document(filename: str, content_type: str, data: bytes, max_spreadsheet_rows: int) -> ParsedDocument:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf" or data.startswith(b"%PDF"):
        return parse_pdf(data)

    office_kind = ""
    if data.startswith(b"PK"):
        try:
            with zipfile.ZipFile(BytesIO(data)) as archive:
                members = set(archive.namelist())
            if "word/document.xml" in members:
                office_kind = "docx"
            elif "xl/workbook.xml" in members:
                office_kind = "xlsx"
            elif "ppt/presentation.xml" in members:
                office_kind = "pptx"
        except zipfile.BadZipFile:
            pass
    if office_kind and ext in {".docx", ".xlsx", ".pptx"} and ext[1:] != office_kind:
        raise ValueError(f"Содержимое файла не соответствует расширению {ext}")
    if office_kind == "docx" or ext == ".docx":
        return parse_docx(data)
    if office_kind == "xlsx" or ext == ".xlsx":
        return parse_xlsx(data, max_spreadsheet_rows)
    if office_kind == "pptx" or ext == ".pptx":
        return parse_pptx(data)
    if ext in {".mht", ".mhtml"} or data[:1024].lstrip().lower().startswith((b"mime-version:", b"from:")):
        return parse_mht(data)
    if ext == ".rtf" or data.lstrip().startswith(b"{\\rtf"):
        return parse_rtf(data)
    if ext in {".html", ".htm"} or "html" in content_type:
        text, _encoding = _decode_text(data)
        return _parse_html(text)
    if ext in {".md", ".markdown"} or content_type == "text/markdown":
        text, _encoding = _decode_text(data)
        return _parse_markdown(text)
    if ext == ".txt" or content_type.startswith("text/plain"):
        text, encoding = _decode_text(data)
        warnings = [] if encoding.startswith("utf-8") else [f"Кодировка текста определена как {encoding}"]
        return ParsedDocument("plaintext", "high" if not warnings else "medium", _paragraph_blocks(text), warnings)
    raise ValueError(f"Unsupported format: {ext or content_type}")
