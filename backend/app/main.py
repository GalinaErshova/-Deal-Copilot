from __future__ import annotations

import hashlib
import io
import json
import math
import secrets
import time
import uuid
from pathlib import Path

from fastapi import Depends, FastAPI, File, Header, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel
from sqlalchemy import inspect, text
from sqlalchemy.orm import Session

from .config import settings
from .db import Base, SessionLocal, engine, get_db
from .document_processing import chunk_extraction_text, parse_document
from .formula_engine import (
    DEFAULT_FORMULAS,
    FORMULA_BY_KEY,
    FORMULAS,
    FormulaError,
    evaluate_formula,
    validate_formula,
)
from .model_gateway import gateway
from .models import (
    Calculation,
    CalculationFormula,
    Deal,
    DealComponentPriceSelection,
    Document,
    ExtractedField,
    ManualServiceLine,
    PipelineRun,
    PipelineStep,
    PriceListItem,
    ReferenceRate,
    utc_now_naive,
)
from .organization_extractor import extract_organization_profile
from .pricing import calculate
from .schemas import (
    CalculationBreakdownRequest,
    CalculationFormulaUpdate,
    CalculationRequest,
    DealComponentPriceSelectionUpdate,
    DealExtraction,
    ManualAreaFieldRequest,
    ManualServiceLineRequest,
    ManualServiceLineUpdate,
    PriceListItemCreate,
    PriceListItemUpdate,
    ProposalExportRequest,
)
from .scope_curator import curate_scope_components
from .scope_extractor import extract_area_components
from .scope_schedule import calculate_monthly_shifts

settings.ensure_dirs()
Base.metadata.create_all(engine)

# Обновляем ранее созданную SQLite-схему: create_all добавляет таблицы, но не колонки.
def ensure_price_list_columns() -> None:
    columns={column["name"] for column in inspect(engine).get_columns("manual_service_lines")}
    with engine.begin() as connection:
        if "price_list_item_id" not in columns:
            connection.execute(text(
                'ALTER TABLE "manual_service_lines" ADD COLUMN "price_list_item_id" '
                'INTEGER REFERENCES "price_list_items" (id) ON DELETE SET NULL'
            ))
        if "price_per_m2_month" not in columns:
            connection.execute(text(
                'ALTER TABLE "manual_service_lines" ADD COLUMN "price_per_m2_month" FLOAT'
            ))

ensure_price_list_columns()

app = FastAPI(title="Deal Copilot API", version="0.1.0")
app.add_middleware(CORSMiddleware, allow_origins=settings.parsed_cors_origins, allow_credentials=True,
                   allow_methods=["*"], allow_headers=["*"])

def seed() -> None:
    db = SessionLocal()
    try:
        if db.query(ReferenceRate).count() == 0:
            for category,name,unit,value,notes in settings.demo_reference_rates:
                db.add(ReferenceRate(category=category,name=name,unit=unit,value=value,notes=notes))
            db.commit()
        existing_formula_keys={row.key for row in db.query(CalculationFormula).all()}
        for key,expression in DEFAULT_FORMULAS.items():
            if key not in existing_formula_keys:
                db.add(CalculationFormula(key=key,expression=expression))
        db.commit()
        # Загружаем демонстрационные услуги только в пустой прайс-лист,
        # чтобы при перезапуске не затронуть введённые пользователем позиции.
        if db.query(PriceListItem).count()==0:
            demo_catalog_path=Path(__file__).resolve().parents[2]/"data"/"demo"/"company-price-list.json"
            if demo_catalog_path.is_file():
                demo_catalog=json.loads(demo_catalog_path.read_text(encoding="utf-8"))
                for item in demo_catalog.get("items",[]):
                    db.add(PriceListItem(**item))
                db.commit()
    finally:
        db.close()
seed()

def _saved_formula_expressions(db:Session)->dict[str,str]:
    return {row.key:row.expression for row in db.query(CalculationFormula).all()}

def _calculation_source_fingerprint(db:Session, deal_id:int)->str:
    """Hash persisted inputs that can change the meaning of a calculation."""
    documents=db.query(Document).filter_by(deal_id=deal_id).order_by(Document.id).all()
    fields=db.query(ExtractedField).filter_by(deal_id=deal_id).order_by(ExtractedField.id).all()
    lines=db.query(ManualServiceLine).filter_by(deal_id=deal_id).order_by(ManualServiceLine.id).all()
    selections=db.query(DealComponentPriceSelection).filter_by(deal_id=deal_id).order_by(
        DealComponentPriceSelection.id).all()
    prices=db.query(PriceListItem).order_by(PriceListItem.id).all()
    rates=db.query(ReferenceRate).order_by(ReferenceRate.id).all()
    source={
        "documents":[(row.id,row.filename,row.parse_status,row.parse_json,row.extracted_text)
                     for row in documents],
        "fields":[(row.id,row.key,row.value,row.unit,row.status,row.confirmed,
                   row.source_document_id,row.source_location,row.source_fragment) for row in fields],
        "manual_lines":[(row.id,row.address,row.area_type,row.work_type,row.area_m2,
                         row.schedule_mode,row.shifts_per_month,row.productivity_m2_per_shift,
                         row.price_list_item_id,row.price_per_m2_month) for row in lines],
        "selections":[(row.component_id,row.price_list_item_id) for row in selections],
        "prices":[(row.id,row.name,row.area_type,row.work_type,row.price_per_m2_month,
                   row.productivity_m2_per_shift,row.is_active) for row in prices],
        "rates":[(row.id,row.category,row.name,row.unit,row.value) for row in rates],
        "formulas":_saved_formula_expressions(db),
        "calculation_settings":{
            "hours_per_shift":settings.hours_per_shift,
            "working_days_per_month":settings.working_days_per_month,
            "working_days_per_week":settings.working_days_per_week,
            "monthly_frequency_shifts":settings.monthly_frequency_shifts,
            "minimum_physical_staff":settings.minimum_physical_staff,
            "currency_decimal_places":settings.currency_decimal_places,
            "labor_hours_decimal_places":settings.labor_hours_decimal_places,
            "fte_decimal_places":settings.fte_decimal_places,
            "sensitivity_deltas":settings.parsed_sensitivity_deltas,
        },
    }
    return hashlib.sha256(json.dumps(source,sort_keys=True,ensure_ascii=False).encode()).hexdigest()

def _formula_payload(db:Session)->list[dict]:
    saved=_saved_formula_expressions(db)
    return [{"key":item.key,"label":item.label,"description":item.description,
             "expression":saved.get(item.key,item.expression),"default_expression":item.expression,
             "variables":list(item.variables)} for item in FORMULAS]

@app.get("/api/formulas")
def calculation_formulas(db:Session=Depends(get_db)):
    return _formula_payload(db)

@app.put("/api/formulas")
def update_calculation_formulas(payload:CalculationFormulaUpdate,db:Session=Depends(get_db)):
    submitted={item.key:item.expression.strip() for item in payload.formulas}
    required=set(FORMULA_BY_KEY)
    if len(submitted)!=len(payload.formulas) or set(submitted)!=required:
        raise HTTPException(422,"Передайте ровно один вариант каждой формулы из списка настроек")
    sample={"area_m2":100.0,"productivity_m2_per_shift":500.0,"hours_per_shift":8.0,
            "working_days_per_month":22.0,"shifts_per_month":22.0,
            "service_price_per_m2_month":180.0,"vat_rate":0.22,"hourly_staff_cost":350.0,
            "replacement_coefficient":1.12,"materials_per_m2_month":7.0,
            "equipment_per_m2_month":2.0,"manager_monthly_cost":8000.0,
            "logistics_monthly":3000.0,"contingency_rate":0.03,"overhead_rate":0.08,
            "target_margin":0.15,"revenue_with_vat":18000.0,"revenue_net":14754.098,
            "labor_hours_month":35.2,"labor_cost":13798.4,"materials_cost":700.0,
            "equipment_cost":200.0,"direct_cost":21698.4,"contingency_cost":650.95,
            "overhead_cost":1180.33,"full_cost":23529.68,"profit":-8775.58,
            "monthly_price":18000.0,"contract_months":10.0}
    try:
        for key,expression in submitted.items():
            definition=FORMULA_BY_KEY[key]
            normalized=validate_formula(expression,definition.variables)
            trial=evaluate_formula(normalized,definition.variables,sample)
            if key in {"labor_hours_month","line_labor_hours_month","revenue_with_vat",
                       "revenue_net","labor_cost","materials_cost","equipment_cost",
                       "direct_cost","contingency_cost","overhead_cost","full_cost",
                       "line_monthly_price","contract_total_price"} and trial<0:
                raise FormulaError(f"{key}: значение не может быть отрицательным")
    except FormulaError as exc:
        raise HTTPException(422,str(exc)) from exc
    existing={row.key:row for row in db.query(CalculationFormula).all()}
    for key,expression in submitted.items():
        if key in existing: existing[key].expression=expression
        else: db.add(CalculationFormula(key=key,expression=expression))
    db.commit()
    return _formula_payload(db)

def _normalize_evidence(text: str | None) -> str:
    return " ".join((text or "").split()).casefold()

def _field_source_is_valid(field, parsed_document: dict) -> bool:
    location = field.source_location or ""
    fragment = _normalize_evidence(field.source_fragment)
    if not location or not fragment:
        return False
    for block in parsed_document.get("blocks", []):
        path = block.get("path", "")
        page = block.get("page_no")
        expected_locations = {path}
        if page:
            expected_locations.add(f"{path} page {page}")
        if location in expected_locations and fragment in _normalize_evidence(block.get("text", "")):
            return True
    return False


def _merge_extractions(extractions: list[DealExtraction]) -> DealExtraction:
    """Объединяет ответы по частям и отмечает разные значения одного поля."""
    fields = []
    missing_fields = []
    contradictions = []
    seen_fields = set()
    seen_missing = set()
    seen_contradictions = set()

    for extraction in extractions:
        for field in extraction.fields:
            identity = (
                field.key,
                field.value,
                field.unit,
                field.source_document,
                field.source_location,
                field.source_fragment,
            )
            if identity not in seen_fields:
                seen_fields.add(identity)
                fields.append(field)
        for missing in extraction.missing_fields:
            if missing not in seen_missing:
                seen_missing.add(missing)
                missing_fields.append(missing)
        for contradiction in extraction.contradictions:
            identity = json.dumps(contradiction, ensure_ascii=False, sort_keys=True)
            if identity not in seen_contradictions:
                seen_contradictions.add(identity)
                contradictions.append(contradiction)

    by_key: dict[str, list] = {}
    for field in fields:
        by_key.setdefault(field.key, []).append(field)
    for key, variants in by_key.items():
        # Группируем по нормализованному значению и единице, чтобы сверить результаты разных частей.
        value_pairs = {
            ((field.value or "").strip().casefold(), (field.unit or "").strip().casefold())
            for field in variants
        }
        if len(value_pairs) > 1:
            contradiction = {
                "key": key,
                "values": [
                    {
                        "value": field.value,
                        "unit": field.unit,
                        "source_document": field.source_document,
                        "source_location": field.source_location,
                        "source_fragment": field.source_fragment,
                    }
                    for field in variants
                ],
            }
            identity = json.dumps(contradiction, ensure_ascii=False, sort_keys=True)
            if identity not in seen_contradictions:
                seen_contradictions.add(identity)
                contradictions.append(contradiction)

    found_keys = set(by_key)
    found_labels = {field.label.casefold() for field in fields}
    missing_fields = [
        item for item in missing_fields
        if item not in found_keys and item.casefold() not in found_labels
    ]
    return DealExtraction(fields=fields, missing_fields=missing_fields, contradictions=contradictions)

class ModelSelectionRequest(BaseModel):
    """Принимает только код заранее настроенного профиля модели."""

    profile_code: str


def require_model_admin(x_model_admin_token: str = Header(default="")) -> None:
    """Разрешает смену модели только владельцу серверного административного токена."""
    if not settings.model_admin_token:
        raise HTTPException(503, "Переключение модели не настроено")
    if not secrets.compare_digest(x_model_admin_token, settings.model_admin_token):
        raise HTTPException(403, "Недостаточно прав для смены модели")


@app.get("/api/admin/model-profiles", dependencies=[Depends(require_model_admin)])
def list_model_profiles(db: Session = Depends(get_db)):
    """Возвращает разрешённые профили и текущий выбор без адресов и секретов."""
    active = gateway.active_profile(db)
    return {
        "active": active.code,
        "profiles": [vars(profile) for profile in gateway.profiles().values()],
    }


@app.put("/api/admin/model-profile", dependencies=[Depends(require_model_admin)])
def update_model_profile(payload: ModelSelectionRequest, db: Session = Depends(get_db)):
    """Меняет активный профиль для следующих запросов без перезапуска API."""
    try:
        profile = gateway.select_profile(db, payload.profile_code)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    return vars(profile)


@app.get("/api/settings")
def public_settings(db: Session = Depends(get_db)):
    active_model = gateway.active_profile(db)
    return {
        "calculation_defaults": settings.calculation_defaults,
        "working_days_per_month": settings.working_days_per_month,
        "working_days_per_week": settings.working_days_per_week,
        "monthly_frequency_shifts": settings.monthly_frequency_shifts,
        "hours_per_shift": settings.hours_per_shift,
        "accepted_upload_extensions": settings.parsed_upload_extensions,
        "demo_mode": active_model.provider == "mock",
        "llm_provider": active_model.provider,
        "llm_model": active_model.model,
        "display_locale": settings.display_locale,
        "currency_code": settings.currency_code,
        "currency_unit_symbol": settings.currency_unit_symbol,
        "display_number_max_fraction_digits": settings.display_number_max_fraction_digits,
        "display_currency_max_fraction_digits": settings.display_currency_max_fraction_digits,
        "display_percentage_factor": settings.display_percentage_factor,
        "display_percentage_decimal_places": settings.display_percentage_decimal_places,
        "base_sensitivity_label": settings.base_sensitivity_label,
        "confidence_good_threshold": settings.confidence_good_threshold,
        "no_bid_margin_threshold": settings.no_bid_margin_threshold,
        "condition_price_decimal_places": settings.condition_price_decimal_places,
        "condition_price_unit": settings.condition_price_unit,
        "sensitivity_bar_min_width": settings.sensitivity_bar_min_width,
        "sensitivity_bar_max_width": settings.sensitivity_bar_max_width,
        "sensitivity_bar_margin_offset": settings.sensitivity_bar_margin_offset,
        "sensitivity_bar_scale": settings.sensitivity_bar_scale,
    }

def add_step(db: Session, run: PipelineRun, name: str, status: str,
             input_data=None, output_data=None, warnings=None, duration_ms=None):
    db.add(PipelineStep(run_id=run.id,name=name,status=status,
                        input_json=json.dumps(input_data,ensure_ascii=False) if input_data is not None else None,
                        output_json=json.dumps(output_data,ensure_ascii=False) if output_data is not None else None,
                        warnings_json=json.dumps(warnings,ensure_ascii=False) if warnings else None,
                        duration_ms=duration_ms))
    db.commit()

@app.get("/api/health")
def health(db: Session = Depends(get_db)):
    active_model = gateway.active_profile(db)
    return {"status":"ok","model":active_model.model,"demo_mode":active_model.provider == "mock"}

@app.post("/api/deals")
def create_deal(title: str = "Новая сделка", db: Session = Depends(get_db)):
    deal = Deal(title=title)
    db.add(deal); db.commit(); db.refresh(deal)
    return {"id":deal.id,"title":deal.title,"status":deal.status}

@app.get("/api/deals")
def list_deals(db: Session = Depends(get_db)):
    deals = db.query(Deal).order_by(Deal.id.desc()).all()
    return [
        {
            "id": deal.id,
            "title": deal.title,
            "status": deal.status,
            "created_at": deal.created_at,
            "document_count": len(deal.documents),
        }
        for deal in deals
    ]

@app.post("/api/deals/{deal_id}/documents")
async def upload_documents(deal_id: int, files: list[UploadFile] = File(...), db: Session = Depends(get_db)):
    deal = db.get(Deal, deal_id)
    if not deal: raise HTTPException(404,"Deal not found")
    if len(files) > settings.max_files_per_upload:
        raise HTTPException(413, f"Не более {settings.max_files_per_upload} файлов за загрузку")

    allowed_extensions = set(settings.parsed_upload_extensions)
    pending_files = []
    total_size = 0
    for upload in files:
        original_name = upload.filename or "file"
        extension = Path(original_name).suffix.lower()
        if extension not in allowed_extensions:
            raise HTTPException(415, f"Формат {extension or '(без расширения)'} не поддерживается")
        data = await upload.read(settings.max_upload_bytes + 1)
        if len(data) > settings.max_upload_bytes:
            raise HTTPException(413, f"Файл {original_name} превышает лимит размера")
        if not data:
            raise HTTPException(400, f"Файл {original_name} пуст")
        total_size += len(data)
        if total_size > settings.max_upload_total_bytes:
            raise HTTPException(413, "Общий размер файлов превышает лимит загрузки")
        pending_files.append((upload, original_name, data))

    run = PipelineRun(deal_id=deal_id); db.add(run); db.commit(); db.refresh(run)
    accepted=[]
    for f, original_name, data in pending_files:
        filename = f"{uuid.uuid4().hex}_{Path(original_name).name}"
        path = settings.resolve_path(settings.upload_dir) / filename
        path.write_bytes(data)
        doc = Document(deal_id=deal_id,filename=original_name,content_type=f.content_type or "",
                       file_path=str(path),parse_status="uploaded")
        db.add(doc); db.commit(); db.refresh(doc)
        accepted.append({"id":doc.id,"filename":doc.filename})
    add_step(db,run,"upload","success",{"files":[x["filename"] for x in accepted]},{"accepted":accepted})
    run.status="success"; run.finished_at=utc_now_naive(); db.commit()
    return {"run_id":run.id,"documents":accepted}

@app.post("/api/deals/{deal_id}/process")
def process_deal(deal_id: int, db: Session = Depends(get_db)):
    deal=db.get(Deal,deal_id)
    if not deal: raise HTTPException(404,"Deal not found")
    docs=db.query(Document).filter(Document.deal_id==deal_id).all()
    if not docs: raise HTTPException(400,"Upload documents first")
    run=PipelineRun(deal_id=deal_id); db.add(run); db.commit(); db.refresh(run)

    combined=[]
    for doc in docs:
        started=time.perf_counter()
        try:
            parsed=parse_document(
                doc.filename,
                doc.content_type,
                settings.resolve_path(doc.file_path).read_bytes(),
                max_spreadsheet_rows=settings.max_spreadsheet_rows,
            )
            doc.parser=parsed.parser; doc.parse_confidence=parsed.confidence
            doc.extracted_text=parsed.plain_text; doc.parse_json=parsed.to_json(); doc.parse_status="parsed"
            db.commit()
            combined.append(parsed.extraction_text(f"{doc.id}:{doc.filename}"))
            add_step(db,run,f"parse:{doc.filename}","success",
                     {"document_id":doc.id},{"parser":parsed.parser,"blocks":len(parsed.blocks),"confidence":parsed.confidence},
                     parsed.warnings,int((time.perf_counter()-started)*1000))
        except Exception as exc:  # noqa: BLE001 — isolate a bad document and continue the batch
            doc.parse_status="failed"; db.commit()
            add_step(db,run,f"parse:{doc.filename}","failed",{"document_id":doc.id},{"error":str(exc)})
    if not combined:
        run.status="failed"; run.finished_at=utc_now_naive(); db.commit()
        raise HTTPException(422,"Не удалось извлечь текст ни из одного документа")
    prompt_text="\n\n".join(combined)
    if len(prompt_text)>settings.max_extraction_chars:
        omitted=len(prompt_text)-settings.max_extraction_chars
        run.status="failed";run.finished_at=utc_now_naive();db.commit()
        add_step(db,run,"ai_extraction","failed",{"documents":len(docs)},
                 {"error":"Лимит текста превышен","omitted_characters":omitted,
                  "document_names":[doc.filename for doc in docs]},
                 [f"Не обработано {omitted} символов; уменьшите набор документов или увеличьте лимит"])
        raise HTTPException(413,
            f"Общий текст документов превышает лимит на {omitted} символов. "
            "Обработка остановлена: уменьшите набор документов или увеличьте лимит.")
    started=time.perf_counter()
    active_model = None
    try:
        active_model = gateway.active_profile(db)
        local_model = active_model.provider in {"local", "ollama"}
        chunks = chunk_extraction_text(prompt_text, settings.llm_input_chunk_chars) if local_model else [prompt_text]
        partial_extractions = [
            gateway.structured(
                task="extract_requirements",
                system=settings.extraction_system_prompt,
                user=chunk,
                schema=DealExtraction,
                profile=active_model,
            )
            for chunk in chunks
        ]
        extraction = _merge_extractions(partial_extractions) if len(partial_extractions) > 1 else partial_extractions[0]
    except Exception as exc:
        run.status="failed"; run.finished_at=utc_now_naive(); db.commit()
        add_step(db,run,"ai_extraction","failed",{"documents":len(docs)},
                 {"error":str(exc),"prompt_version":settings.extraction_prompt_version,
                  "model_profile":active_model.code if active_model else None})
        raise HTTPException(502,"Не удалось обработать документы моделью") from exc
    db.query(ExtractedField).filter(ExtractedField.deal_id==deal_id).delete()
    document_by_marker={f"{doc.id}:{doc.filename}":doc for doc in docs}
    parsed_by_document_id = {
        doc.id: json.loads(doc.parse_json) if doc.parse_json else {"blocks": []}
        for doc in docs
    }
    for field in extraction.fields:
        source_doc = document_by_marker.get(field.source_document or "")
        source_verified = bool(
            source_doc
            and _field_source_is_valid(
                field,
                parsed_by_document_id.get(source_doc.id, {"blocks": []}),
            )
        )
        db.add(ExtractedField(deal_id=deal_id,key=field.key,label=field.label,value=field.value,
                              unit=field.unit,status="source_verified" if source_verified else "source_unverified",
                              confidence=field.confidence,
                              source_document_id=source_doc.id if source_doc else None,
                              source_location=field.source_location,source_fragment=field.source_fragment))
    db.commit()
    add_step(db,run,"ai_extraction","success",{"documents":len(docs), "chunks":len(chunks)},
             {**extraction.model_dump(),"prompt_version":settings.extraction_prompt_version,
              "model_profile":active_model.code,"model":active_model.model},
             duration_ms=int((time.perf_counter()-started)*1000))
    run.status="success"; run.finished_at=utc_now_naive(); db.commit()
    return extraction

@app.get("/api/deals/{deal_id}/documents")
def documents(deal_id:int, db: Session=Depends(get_db)):
    docs=db.query(Document).filter(Document.deal_id==deal_id).all()
    return [{"id":d.id,"filename":d.filename,"parser":d.parser,"confidence":d.parse_confidence,
             "status":d.parse_status,"warnings":(json.loads(d.parse_json).get("warnings",[]) if d.parse_json else [])} for d in docs]

@app.get("/api/deals/{deal_id}/organization-profile")
def organization_profile(deal_id:int, db:Session=Depends(get_db)):
    """Ищет реквизиты заказчика в уже разобранных документах и возвращает ссылки на источники."""
    if not db.get(Deal,deal_id): raise HTTPException(404,"Deal not found")
    docs=db.query(Document).filter(Document.deal_id==deal_id).order_by(Document.id).all()
    return extract_organization_profile(docs)

@app.get("/api/documents/{document_id}/original")
def document_original(document_id:int, db:Session=Depends(get_db)):
    doc=db.get(Document,document_id)
    if not doc: raise HTTPException(404,"Document not found")
    return FileResponse(settings.resolve_path(doc.file_path),media_type=doc.content_type or "application/octet-stream",filename=doc.filename)

@app.get("/api/documents/{document_id}/parsed")
def document_parsed(document_id:int, db:Session=Depends(get_db)):
    doc=db.get(Document,document_id)
    if not doc: raise HTTPException(404,"Document not found")
    return json.loads(doc.parse_json) if doc.parse_json else {"blocks":[],"warnings":[]}

@app.get("/api/deals/{deal_id}/fields")
def fields(deal_id:int,db:Session=Depends(get_db)):
    rows=db.query(ExtractedField).filter(ExtractedField.deal_id==deal_id).all()
    return [{"id":r.id,"key":r.key,"label":r.label,"value":r.value,"unit":r.unit,"status":r.status,
             "confidence":r.confidence,"source_document_id":r.source_document_id,
             "source_location":r.source_location,"source_fragment":r.source_fragment,"confirmed":r.confirmed} for r in rows]

def _productivity_reference(component:dict, rates:list[ReferenceRate])->ReferenceRate|None:
    normalized_type=str(component.get("work_type") or "").casefold()
    is_snow="снег" in normalized_type or "механизирован" in normalized_type
    for rate in rates:
        name=str(rate.name or "").casefold()
        unit=str(rate.unit or "").casefold().replace(" ", "")
        if unit not in {"м²/смену", "м2/смену"} or rate.value <= 0:
            continue
        rate_is_snow=any(token in name for token in ("снег", "территор", "механизирован"))
        if rate_is_snow == is_snow and (is_snow or "регуляр" in name or "помещ" in name):
            return rate
    return None

@app.get("/api/deals/{deal_id}/area-components")
def area_components(deal_id:int, db:Session=Depends(get_db)):
    if not db.get(Deal, deal_id): raise HTTPException(404,"Deal not found")
    docs=db.query(Document).filter(Document.deal_id==deal_id).order_by(Document.id).all()
    components=extract_area_components(docs)
    components=curate_scope_components(docs,components)
    rates=db.query(ReferenceRate).filter(ReferenceRate.category=="productivity").all()
    price_items=db.query(PriceListItem).filter(PriceListItem.is_active.is_(True)).all()
    price_by_work_type={" ".join(item.work_type.casefold().split()):item for item in price_items}
    saved_price_selections={selection.component_id:selection for selection in db.query(DealComponentPriceSelection).filter(
        DealComponentPriceSelection.deal_id==deal_id
    ).all()}
    for component in components:
        rate=_productivity_reference(component,rates)
        saved_selection=saved_price_selections.get(component["id"])
        if saved_selection:
            price_item=db.get(PriceListItem,saved_selection.price_list_item_id) if saved_selection.price_list_item_id else None
        else:
            price_item=price_by_work_type.get(" ".join(str(component.get("work_type") or "").casefold().split()))
        component["productivity_m2_per_shift"]=(
            price_item.productivity_m2_per_shift if price_item and price_item.productivity_m2_per_shift
            else rate.value if rate else None
        )
        component["productivity_reference"]=(
            {"id":rate.id,"name":rate.name,"unit":rate.unit,"value":rate.value,"notes":rate.notes}
            if rate else None
        )
        if price_item:
            component["price_list_item_id"]=price_item.id
            component["price_per_m2_month"]=price_item.price_per_m2_month
            component["company_service_name"]=price_item.name
        elif saved_selection:
            component["price_list_item_id"]=None
            component["price_per_m2_month"]=None
        component["shifts_per_month"]=calculate_monthly_shifts(
            str(component.get("schedule_mode") or "unspecified"),
            working_days_per_month=settings.working_days_per_month,
            working_days_per_week=settings.working_days_per_week,
            monthly_frequency_shifts=settings.monthly_frequency_shifts,
        )
    manual_lines=db.query(ManualServiceLine).filter(ManualServiceLine.deal_id==deal_id).order_by(ManualServiceLine.id).all()
    for line in manual_lines:
        component={
            "id":f"manual:{line.id}","origin":"manual","manual_line_id":line.id,
            "address":line.address,"area_type":line.area_type,"work_type":line.work_type,
            "work_type_source_document_id":None,"work_type_source_location":"","work_type_source_fragment":"",
            "price_list_item_id":line.price_list_item_id,"price_per_m2_month":line.price_per_m2_month,
            "area_m2":line.area_m2,"source_document_id":None,"source_document_name":"",
            "source_location":"","source_fragment":"","schedule_mode":line.schedule_mode,
            "schedule_label":{"daily":"Каждый рабочий день","weekly":"Еженедельно","monthly":"Ежемесячно",
                              "on_request":"По заявкам","custom":"Другой режим","unspecified":"Не указан"}.get(line.schedule_mode,line.schedule_mode),
            "schedule_status":"needs_input" if line.schedule_mode=="unspecified" else "verified",
            "schedule_warnings":[],"schedule_source_document_id":None,"schedule_source_document_name":None,
            "schedule_source_location":"","schedule_source_fragment":"","schedule_additional_frequencies":[],
            "curation_status":"verified","curation_warnings":[],
            "productivity_m2_per_shift":line.productivity_m2_per_shift,"productivity_reference":None,
            "shifts_per_month":calculate_monthly_shifts(
                line.schedule_mode,working_days_per_month=settings.working_days_per_month,
                working_days_per_week=settings.working_days_per_week,
                monthly_frequency_shifts=settings.monthly_frequency_shifts,manual_shifts=line.shifts_per_month),
        }
        rate=_productivity_reference(component,rates)
        if component["productivity_m2_per_shift"] is None and rate:
            component["productivity_m2_per_shift"]=rate.value
        component["productivity_reference"]=(
            {"id":rate.id,"name":rate.name,"unit":rate.unit,"value":rate.value,"notes":rate.notes}
            if rate else None
        )
        components.append(component)
    return components

@app.put("/api/deals/{deal_id}/component-price-selection")
def save_component_price_selection(
    deal_id:int,payload:DealComponentPriceSelectionUpdate,db:Session=Depends(get_db)
):
    if not db.get(Deal,deal_id): raise HTTPException(404,"Сделка не найдена")
    documents=db.query(Document).filter(Document.deal_id==deal_id).order_by(Document.id).all()
    document_components=extract_area_components(documents)
    if not any(component["id"]==payload.component_id for component in document_components):
        raise HTTPException(404,"Строка требований не найдена в документах сделки")
    if payload.price_list_item_id is not None:
        item=db.get(PriceListItem,payload.price_list_item_id)
        if not item or not item.is_active: raise HTTPException(404,"Активная услуга в прайс-листе не найдена")
    selection=db.query(DealComponentPriceSelection).filter_by(
        deal_id=deal_id,component_id=payload.component_id
    ).first()
    if selection:
        selection.price_list_item_id=payload.price_list_item_id
    else:
        selection=DealComponentPriceSelection(
            deal_id=deal_id,component_id=payload.component_id,
            price_list_item_id=payload.price_list_item_id,
        )
        db.add(selection)
    db.commit()
    return {"component_id":selection.component_id,"price_list_item_id":selection.price_list_item_id}

@app.post("/api/deals/{deal_id}/area-components/manual")
def create_manual_service_line(deal_id:int,payload:ManualServiceLineRequest,db:Session=Depends(get_db)):
    if not db.get(Deal,deal_id): raise HTTPException(404,"Deal not found")
    values=payload.model_dump()
    price_item=db.get(PriceListItem,values["price_list_item_id"]) if values["price_list_item_id"] else None
    if values["price_list_item_id"] and (not price_item or not price_item.is_active):
        raise HTTPException(422,"Выберите активную позицию из прайс-листа")
    if price_item:
        values["price_per_m2_month"]=values["price_per_m2_month"] or price_item.price_per_m2_month
        values["productivity_m2_per_shift"]=values["productivity_m2_per_shift"] or price_item.productivity_m2_per_shift
    line=ManualServiceLine(deal_id=deal_id,**values)
    db.add(line);db.commit();db.refresh(line)
    return {"id":line.id,"component_id":f"manual:{line.id}","origin":"manual"}

@app.patch("/api/deals/{deal_id}/area-components/manual/{line_id}")
def update_manual_service_line(deal_id:int,line_id:int,payload:ManualServiceLineUpdate,db:Session=Depends(get_db)):
    line=db.query(ManualServiceLine).filter(ManualServiceLine.id==line_id,ManualServiceLine.deal_id==deal_id).first()
    if not line: raise HTTPException(404,"Ручная строка услуг не найдена")
    updates=payload.model_dump(exclude_unset=True)
    if updates.get("price_list_item_id"):
        price_item=db.get(PriceListItem,updates["price_list_item_id"])
        if not price_item or not price_item.is_active:
            raise HTTPException(422,"Выберите активную позицию из прайс-листа")
        updates.setdefault("price_per_m2_month",price_item.price_per_m2_month)
        updates.setdefault("productivity_m2_per_shift",price_item.productivity_m2_per_shift)
    for key,value in updates.items(): setattr(line,key,value)
    db.commit()
    return {"id":line.id,"component_id":f"manual:{line.id}","origin":"manual"}

@app.delete("/api/deals/{deal_id}/area-components/manual/{line_id}")
def delete_manual_service_line(deal_id:int,line_id:int,db:Session=Depends(get_db)):
    line=db.query(ManualServiceLine).filter(ManualServiceLine.id==line_id,ManualServiceLine.deal_id==deal_id).first()
    if not line: raise HTTPException(404,"Ручная строка услуг не найдена")
    db.delete(line);db.commit()
    return {"ok":True}

@app.patch("/api/fields/{field_id}")
def update_field(field_id:int,payload:dict,db:Session=Depends(get_db)):
    row=db.get(ExtractedField,field_id)
    if not row: raise HTTPException(404,"Field not found")
    if "value" in payload: row.value=str(payload["value"]) if payload["value"] is not None else None
    row.confirmed=bool(payload.get("confirmed",True))
    row.status="user_changed" if "value" in payload else "user_confirmed"
    db.commit()
    return {"ok":True}

@app.post("/api/deals/{deal_id}/fields/area")
def create_manual_area_field(deal_id:int,payload:ManualAreaFieldRequest,db:Session=Depends(get_db)):
    if not db.get(Deal,deal_id): raise HTTPException(404,"Deal not found")
    row=ExtractedField(deal_id=deal_id,key="area_m2",label="Площадь",value=str(payload.value),unit="м²",
                       status="manual_unverified",confirmed=False)
    db.add(row); db.commit(); db.refresh(row)
    return {"id":row.id,"key":row.key,"value":row.value,"confirmed":row.confirmed}

@app.get("/api/deals/{deal_id}/pipeline")
def pipeline(deal_id:int,db:Session=Depends(get_db)):
    runs=db.query(PipelineRun).filter(PipelineRun.deal_id==deal_id).order_by(PipelineRun.id.desc()).all()
    result=[]
    for run in runs:
        steps=db.query(PipelineStep).filter(PipelineStep.run_id==run.id).order_by(PipelineStep.id).all()
        result.append({"id":run.id,"status":run.status,"started_at":run.started_at,"finished_at":run.finished_at,
            "steps":[{"name":s.name,"status":s.status,"duration_ms":s.duration_ms,
                      "input":json.loads(s.input_json) if s.input_json else None,
                      "output":json.loads(s.output_json) if s.output_json else None,
                      "warnings":json.loads(s.warnings_json) if s.warnings_json else []} for s in steps]})
    return result

@app.get("/api/reference-rates")
def reference_rates(db:Session=Depends(get_db)):
    return [{"id":r.id,"category":r.category,"name":r.name,"unit":r.unit,"value":r.value,"notes":r.notes}
            for r in db.query(ReferenceRate).all()]

def _price_list_payload(item:PriceListItem)->dict:
    return {"id":item.id,"name":item.name,"area_type":item.area_type,"work_type":item.work_type,
            "price_per_m2_month":item.price_per_m2_month,
            "productivity_m2_per_shift":item.productivity_m2_per_shift,
            "notes":item.notes,"is_active":item.is_active}

@app.get("/api/price-list")
def company_price_list(db:Session=Depends(get_db)):
    return [_price_list_payload(item) for item in db.query(PriceListItem).order_by(PriceListItem.name,PriceListItem.id).all()]

@app.get("/api/demo-provider-tariffs")
def demo_provider_tariffs():
    """Справочные демо-тарифы исполнителей с исходными единицами и ссылками."""
    path=Path(__file__).resolve().parents[2]/"data"/"demo"/"provider-price-rates.json"
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise HTTPException(404,"Демо-каталог тарифов не найден") from exc

@app.post("/api/price-list")
def create_price_list_item(payload:PriceListItemCreate,db:Session=Depends(get_db)):
    item=PriceListItem(**payload.model_dump())
    db.add(item);db.commit();db.refresh(item)
    return _price_list_payload(item)

@app.patch("/api/price-list/{item_id}")
def update_price_list_item(item_id:int,payload:PriceListItemUpdate,db:Session=Depends(get_db)):
    item=db.get(PriceListItem,item_id)
    if not item: raise HTTPException(404,"Позиция прайс-листа не найдена")
    for key,value in payload.model_dump(exclude_unset=True).items(): setattr(item,key,value)
    db.commit();db.refresh(item)
    return _price_list_payload(item)

@app.delete("/api/price-list/{item_id}")
def archive_price_list_item(item_id:int,db:Session=Depends(get_db)):
    item=db.get(PriceListItem,item_id)
    if not item: raise HTTPException(404,"Позиция прайс-листа не найдена")
    item.is_active=False
    db.commit()
    return {"ok":True,"is_active":False}

@app.post("/api/deals/{deal_id}/calculate")
def calculate_deal(deal_id:int,req:CalculationRequest,db:Session=Depends(get_db)):
    if not db.get(Deal,deal_id): raise HTTPException(404,"Deal not found")
    confirmed_area = db.query(ExtractedField).filter(
        ExtractedField.deal_id==deal_id,
        ExtractedField.key=="area_m2",
        ExtractedField.confirmed.is_(True),
    ).first()
    if not confirmed_area or not confirmed_area.value:
        raise HTTPException(409,"Подтвердите площадь объекта перед расчётом")
    try:
        confirmed_area_value = float(confirmed_area.value.replace(",", "."))
    except ValueError:
        raise HTTPException(409,"Подтверждённая площадь указана некорректно")
    if not math.isfinite(confirmed_area_value) or confirmed_area_value <= 0:
        raise HTTPException(409,"Подтверждённая площадь должна быть больше нуля")
    if not math.isclose(req.area_m2, confirmed_area_value, rel_tol=settings.area_match_tolerance):
        raise HTTPException(409,"Площадь расчёта отличается от подтверждённой площади")
    run=PipelineRun(deal_id=deal_id); db.add(run); db.commit(); db.refresh(run)
    started=time.perf_counter()
    formula_expressions=_saved_formula_expressions(db)
    try:
        result=calculate(req, settings, formula_expressions)
    except FormulaError as exc:
        raise HTTPException(422,str(exc)) from exc
    result_payload={**result.model_dump(),"formulas_used":formula_expressions,
                    "source_fingerprint":_calculation_source_fingerprint(db,deal_id)}
    calc=Calculation(deal_id=deal_id,input_json=req.model_dump_json(),
                     output_json=json.dumps(result_payload,ensure_ascii=False),decision=result.decision)
    db.add(calc)
    deal=db.get(Deal,deal_id); deal.status="calculated"; db.commit(); db.refresh(calc)
    add_step(
        db,run,"deterministic_calculation","success",
        req.model_dump(),
        {"calculation_id":calc.id,**result_payload},
        duration_ms=int((time.perf_counter()-started)*1000),
    )
    run.status="success"; run.finished_at=utc_now_naive(); db.commit()
    return {"calculation_id":calc.id,**result_payload}

@app.post("/api/deals/{deal_id}/calculate-breakdown")
def calculate_deal_breakdown(deal_id:int,req:CalculationBreakdownRequest,db:Session=Depends(get_db)):
    deal=db.get(Deal,deal_id)
    if not deal: raise HTTPException(404,"Deal not found")
    docs=db.query(Document).filter(Document.deal_id==deal_id).order_by(Document.id).all()
    document_components=curate_scope_components(docs,extract_area_components(docs))
    manual_rows=db.query(ManualServiceLine).filter(ManualServiceLine.deal_id==deal_id).order_by(ManualServiceLine.id).all()
    manual_components=[{
        "id":f"manual:{line.id}","origin":"manual","manual_line_id":line.id,
        "address":line.address,"area_type":line.area_type,"work_type":line.work_type,
        "price_list_item_id":line.price_list_item_id,"price_per_m2_month":line.price_per_m2_month,
        "source_document_id":None,"source_document_name":"","source_location":"","source_fragment":"",
        "schedule_source_document_id":None,"schedule_source_document_name":None,
        "schedule_source_location":"","schedule_source_fragment":"",
    } for line in manual_rows]
    source_components=document_components+manual_components
    saved_selections={selection.component_id:selection for selection in db.query(DealComponentPriceSelection).filter(
        DealComponentPriceSelection.deal_id==deal_id
    ).all()}
    active_price_items={item.id:item for item in db.query(PriceListItem).filter(PriceListItem.is_active.is_(True)).all()}
    price_by_work_type={" ".join(item.work_type.casefold().split()):item for item in active_price_items.values()}
    for source in source_components:
        selection=saved_selections.get(source["id"])
        selected_id=selection.price_list_item_id if selection else source.get("price_list_item_id")
        catalog_item=active_price_items.get(selected_id) if selected_id else None
        if not selection and not selected_id:
            catalog_item=price_by_work_type.get(" ".join(str(source.get("work_type") or "").casefold().split()))
        if catalog_item:
            source["price_list_item_id"]=catalog_item.id
            source["price_per_m2_month"]=catalog_item.price_per_m2_month
            source["company_service_name"]=catalog_item.name
    source_by_id={component["id"]:component for component in source_components}
    requested_ids={component.id for component in req.components}
    if not source_by_id:
        raise HTTPException(409,"Добавьте в требования хотя бы одну услугу из документов или вручную")
    if requested_ids!=set(source_by_id):
        raise HTTPException(409,"Состав адресов и площадей изменился. Обновите страницу и проверьте источники")
    if any(component["curation_status"]!="verified" for component in document_components):
        raise HTTPException(409,"Куратор обнаружил непроверенные источники или неоднозначные строки. Исправьте данные документов перед расчётом")
    if any(component["schedule_status"]=="needs_review" for component in document_components):
        raise HTTPException(409,"Куратор не смог подтвердить источник режима уборки; проверьте цитаты в документе")
    if not req.confirmed:
        raise HTTPException(409,"Подтвердите расчёт целиком после проверки исходных данных и предположений")
    confirmed_term=db.query(ExtractedField).filter(
        ExtractedField.deal_id==deal_id,ExtractedField.key=="contract_months",
        ExtractedField.confirmed.is_(True)).order_by(ExtractedField.id).first()
    if confirmed_term and confirmed_term.value:
        try:
            term=int(confirmed_term.value)
        except (TypeError,ValueError) as exc:
            raise HTTPException(409,"Подтверждённый срок договора указан некорректно") from exc
        if term!=req.assumptions.contract_months:
            raise HTTPException(409,"Срок расчёта отличается от подтверждённого срока договора")

    items_by_id={component.id:component for component in req.components}
    shifts_by_id: dict[str, float] = {}
    manual_by_id={f"manual:{line.id}":line for line in manual_rows}
    for source in source_components:
        item=items_by_id[source["id"]]
        if item.productivity_m2_per_shift is None:
            raise HTTPException(409,f"Задайте выработку для строки «{source['address']} — {source['work_type']}»")
        shifts=calculate_monthly_shifts(
            item.schedule_mode,
            working_days_per_month=settings.working_days_per_month,
            working_days_per_week=settings.working_days_per_week,
            monthly_frequency_shifts=settings.monthly_frequency_shifts,
            manual_shifts=item.shifts_per_month,
        )
        if shifts is None or shifts<=0:
            raise HTTPException(409,f"Укажите ожидаемое число смен в месяц для строки «{source['address']} — {source['work_type']}»")
        shifts_by_id[source["id"]]=shifts

    total_area=sum(items_by_id[component["id"]].area_m2 for component in source_components)
    price_by_id={}
    for source in source_components:
        component_id=source["id"]
        item=items_by_id[component_id]
        if source.get("origin")!="manual" and source.get("price_list_item_id"):
            price_by_id[component_id]=source["price_per_m2_month"]
        else:
            price_by_id[component_id]=(item.price_per_m2_month or source.get("price_per_m2_month")
                                       or req.assumptions.service_price_per_m2_month)
    total_workload=sum(
        items_by_id[component["id"]].area_m2
        / items_by_id[component["id"]].productivity_m2_per_shift
        * shifts_by_id[component["id"]]
        for component in source_components
    )
    if total_area<=0 or total_workload<=0:
        raise HTTPException(409,"Для расчёта нужна положительная площадь и производительность")

    line_results=[]
    total_labor_hours=0.0
    total_monthly_price=0.0
    schedule_labels={"daily":"Каждый рабочий день","weekly":"Еженедельно",
                     "monthly":"Ежемесячно","on_request":"По заявкам",
                     "custom":"Другой режим","unspecified":"Не указан"}
    productivity_rates=db.query(ReferenceRate).filter(ReferenceRate.category=="productivity").all()
    formula_expressions=_saved_formula_expressions(db)
    for source in source_components:
        item=items_by_id[source["id"]]
        if source.get("origin")=="manual":
            line=manual_by_id[source["id"]]
            source.update({"area_type":line.area_type,"curation_status":"verified","curation_warnings":[],
                           "schedule_mode":item.schedule_mode,
                           "schedule_label":schedule_labels.get(item.schedule_mode,item.schedule_mode),
                           "schedule_status":"verified","schedule_warnings":[],"schedule_additional_frequencies":[],
                           "work_type_source_document_id":None,"work_type_source_location":"","work_type_source_fragment":""})
        rate=_productivity_reference(source,productivity_rates)
        shifts=shifts_by_id[source["id"]]
        try:
            labor_hours=evaluate_formula(
                formula_expressions["line_labor_hours_month"],FORMULA_BY_KEY["line_labor_hours_month"].variables,
                {"area_m2":item.area_m2,"productivity_m2_per_shift":item.productivity_m2_per_shift,
                 "hours_per_shift":settings.hours_per_shift,"shifts_per_month":shifts},
            )
        except FormulaError as exc:
            raise HTTPException(422,str(exc)) from exc
        fte=labor_hours/req.assumptions.monthly_hours_per_fte
        try:
            line_monthly_price=evaluate_formula(
                formula_expressions["line_monthly_price"],FORMULA_BY_KEY["line_monthly_price"].variables,
                {"area_m2":item.area_m2,"service_price_per_m2_month":price_by_id[source["id"]]},
            )
            line_contract_price=evaluate_formula(
                formula_expressions["contract_total_price"],FORMULA_BY_KEY["contract_total_price"].variables,
                {"monthly_price":line_monthly_price,"contract_months":req.assumptions.contract_months},
            )
        except FormulaError as exc:
            raise HTTPException(422,str(exc)) from exc
        if labor_hours<0 or line_monthly_price<0 or line_contract_price<0:
            raise HTTPException(422,"Трудозатраты и цена строки не могут быть отрицательными")
        total_labor_hours+=labor_hours
        total_monthly_price+=round(line_monthly_price,settings.currency_decimal_places)
        original_schedule=source.get("schedule_mode")
        schedule_label=(source.get("schedule_label") if original_schedule==item.schedule_mode
                        else schedule_labels.get(item.schedule_mode,item.schedule_mode))
        line_results.append({
            **source,
            "original_schedule_label":source.get("schedule_label"),
            "schedule_label":schedule_label,
            "schedule_manually_changed":original_schedule!=item.schedule_mode,
            "area_m2":item.area_m2,
            "productivity_m2_per_shift":item.productivity_m2_per_shift,
            "price_per_m2_month":price_by_id[source["id"]],
            "schedule_mode":item.schedule_mode,
            "shifts_per_month":shifts,
            "productivity_reference":({"id":rate.id,"name":rate.name,"unit":rate.unit,"notes":rate.notes} if rate else None),
            "labor_hours_month":round(labor_hours,settings.labor_hours_decimal_places),
            "monthly_price":round(line_monthly_price,settings.currency_decimal_places),
            "contract_price":round(line_contract_price,settings.currency_decimal_places),
            "fte":round(fte,settings.fte_decimal_places),
            "physical_staff":max(settings.minimum_physical_staff,math.ceil(fte*req.assumptions.replacement_coefficient)),
        })

    effective_productivity=total_area*settings.working_days_per_month/total_workload
    effective_price=total_monthly_price/total_area
    calculation_request=req.assumptions.model_copy(update={
        "area_m2":total_area,
        "productivity_m2_per_shift":effective_productivity,
        "service_price_per_m2_month":effective_price,
    })
    try:
        aggregate=calculate(calculation_request,settings,formula_expressions,
                            labor_hours_override=total_labor_hours,
                            revenue_with_vat_override=total_monthly_price)
    except FormulaError as exc:
        raise HTTPException(422,str(exc)) from exc
    output={
        "confirmed":req.confirmed,
        "components":line_results,
        "formulas_used":formula_expressions,
        "source_fingerprint":_calculation_source_fingerprint(db,deal_id),
        "total_area_m2":round(total_area,settings.currency_decimal_places),
        "physical_staff_by_site":sum(item["physical_staff"] for item in line_results),
        "aggregate":aggregate.model_dump(),
    }
    run=PipelineRun(deal_id=deal_id);db.add(run);db.commit();db.refresh(run)
    calc=Calculation(deal_id=deal_id,input_json=req.model_dump_json(),
                     output_json=json.dumps(output,ensure_ascii=False),decision=aggregate.decision)
    db.add(calc);deal.status="calculated";db.commit();db.refresh(calc)
    add_step(db,run,"deterministic_calculation_by_area","success",
             req.model_dump(),{"calculation_id":calc.id,**output})
    run.status="success";run.finished_at=utc_now_naive();db.commit()
    return {"calculation_id":calc.id,**aggregate.model_dump(),**output}

@app.get("/api/deals/{deal_id}/calculations")
def calculations(deal_id:int,db:Session=Depends(get_db)):
    rows=db.query(Calculation).filter(Calculation.deal_id==deal_id).order_by(Calculation.id.desc()).all()
    current_fingerprint=_calculation_source_fingerprint(db,deal_id)
    return [{"id":r.id,"created_at":r.created_at,"decision":r.decision,
              "input":json.loads(r.input_json),"output":json.loads(r.output_json),
              "is_current":json.loads(r.output_json).get("source_fingerprint")==current_fingerprint}
            for r in rows]

@app.post("/api/deals/{deal_id}/proposal")
def export_proposal(deal_id:int,req:ProposalExportRequest,db:Session=Depends(get_db)):
    """Формирует редактируемый проект КП из подтверждённых требований и последнего расчёта."""
    deal=db.get(Deal,deal_id)
    if not deal:
        raise HTTPException(404,"Deal not found")
    calculation=db.query(Calculation).filter(Calculation.deal_id==deal_id).order_by(Calculation.id.desc()).first()
    if not calculation:
        raise HTTPException(409,"Сначала выполните расчёт экономики сделки")

    from docx import Document as WordDocument
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    output=json.loads(calculation.output_json)
    if output.get("source_fingerprint")!=_calculation_source_fingerprint(db,deal_id):
        raise HTTPException(409,"Исходные данные изменились. Выполните и подтвердите новый расчёт")
    if output.get("components") and not output.get("confirmed"):
        raise HTTPException(409,"Подтвердите актуальный расчёт перед выгрузкой КП")
    aggregate=output.get("aggregate",output)
    components=output.get("components",[])
    calculation_input=json.loads(calculation.input_json)
    assumptions=calculation_input.get("assumptions",calculation_input)
    price_per_m2=float(assumptions.get("service_price_per_m2_month",0))
    formula_expressions=output.get("formulas_used") or _saved_formula_expressions(db)
    confirmed_fields=db.query(ExtractedField).filter(
        ExtractedField.deal_id==deal_id,ExtractedField.confirmed.is_(True)
    ).order_by(ExtractedField.id).all()
    contract_months=int(assumptions.get("contract_months",1))
    confirmed_term=next((field.value for field in confirmed_fields if field.key=="contract_months" and field.value),None)
    if confirmed_term:
        try:
            if int(confirmed_term)!=contract_months:
                raise HTTPException(409,"Подтверждённый срок отличается от сохранённого расчёта")
        except (TypeError,ValueError) as exc:
            raise HTTPException(409,"Подтверждённый срок договора указан некорректно") from exc
    doc=WordDocument()
    title=doc.add_heading("Коммерческое предложение",0)
    title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Заказчик: {req.customer_name.strip() or '[указать заказчика]'}")
    for label,value in [
        ("Юридический адрес",req.customer_address),
        ("ИНН",req.customer_inn),
        ("КПП",req.customer_kpp),
        ("ОГРН",req.customer_ogrn),
        ("Представитель заказчика",req.customer_contact_person),
        ("Телефон заказчика",req.customer_phone),
        ("Электронная почта заказчика",req.customer_email),
    ]:
        if value.strip(): doc.add_paragraph(f"{label}: {value.strip()}")
    doc.add_paragraph(f"Исполнитель: {req.supplier_name.strip() or '[указать исполнителя]'}")
    if req.contact_details.strip():
        doc.add_paragraph(f"Контакты исполнителя: {req.contact_details.strip()}")
    doc.add_paragraph(f"Предложение действительно {req.validity_days} календарных дней с даты подготовки.")
    doc.add_heading("Предмет предложения",level=1)
    doc.add_paragraph("Оказание услуг по уборке объектов на условиях и в объёме, указанных ниже.")

    if components:
        doc.add_heading("Объекты и виды работ",level=1)
        table=doc.add_table(rows=1,cols=7)
        table.style="Light Shading Accent 1"
        for cell,value in zip(table.rows[0].cells,["Адрес","Вид работ","Площадь, м²","Режим","Тариф, ₽/м²·мес.","Стоимость, ₽/мес.","Стоимость за срок, ₽"]):
            cell.text=value
        for item in components:
            row=table.add_row().cells
            area=float(item.get("area_m2",0) or 0)
            item_price=float(item.get("price_per_m2_month") or price_per_m2)
            monthly_value=float(item.get("monthly_price") or evaluate_formula(
                formula_expressions["line_monthly_price"],FORMULA_BY_KEY["line_monthly_price"].variables,
                {"area_m2":area,"service_price_per_m2_month":item_price}))
            contract_value=float(item.get("contract_price") or evaluate_formula(
                formula_expressions["contract_total_price"],FORMULA_BY_KEY["contract_total_price"].variables,
                {"monthly_price":monthly_value,"contract_months":contract_months}))
            rate=f"{item_price:,.2f}".replace(","," ").replace(".",",")
            monthly=f"{monthly_value:,.2f}".replace(","," ").replace(".",",")
            line_contract=f"{contract_value:,.2f}".replace(","," ").replace(".",",")
            for cell,value in zip(row,[item.get("address","[уточнить]"),item.get("company_service_name") or item.get("work_type","[уточнить]"),
                str(item.get("area_m2","—")),item.get("schedule_label",item.get("schedule_mode","[уточнить]")),rate,monthly,line_contract]):
                cell.text=str(value)

    if confirmed_fields:
        doc.add_heading("Подтверждённые условия",level=1)
        for field in confirmed_fields:
            if field.key=="area_m2" and components:
                continue
            doc.add_paragraph(f"{field.label}: {field.value or '—'} {field.unit or ''}".strip(),style="List Bullet")

    doc.add_heading("Стоимость услуг",level=1)
    monthly_price=f"{aggregate.get('revenue_with_vat',0):,.2f}".replace(","," ").replace(".",",")
    doc.add_paragraph(f"Стоимость за месяц: {monthly_price} ₽, включая НДС при его применении.")
    doc.add_paragraph(f"Срок оказания услуг: {contract_months} мес.")
    if components:
        contract_total=sum(float(item.get("contract_price") or 0) for item in components)
    else:
        contract_total=evaluate_formula(
            formula_expressions["contract_total_price"],FORMULA_BY_KEY["contract_total_price"].variables,
            {"monthly_price":float(aggregate.get("revenue_with_vat",0)),"contract_months":contract_months},
        )
    contract_price=f"{contract_total:,.2f}".replace(","," ").replace(".",",")
    doc.add_paragraph(f"Общая стоимость за срок договора: {contract_price} ₽.")
    doc.add_paragraph("Налоговый режим и ставка НДС подлежат проверке и уточнению перед отправкой.")
    if req.additional_terms.strip():
        doc.add_heading("Дополнительные условия",level=1)
        doc.add_paragraph(req.additional_terms.strip())
    doc.add_paragraph("Проект сформирован по подтверждённым данным сделки. Проверьте реквизиты, налогообложение и договорные условия перед отправкой заказчику.")

    buffer=io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return StreamingResponse(buffer,media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition":f'attachment; filename="commercial-proposal-deal-{deal_id}.docx"'})
