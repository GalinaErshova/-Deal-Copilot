from io import BytesIO

from openpyxl import Workbook

from app.document_processing import ParsedBlock, ParsedDocument, chunk_extraction_text
from app.main import _field_source_is_valid
from app.schemas import DealExtraction, FieldEvidence


def test_extraction_text_contains_traceable_markers():
    parsed = ParsedDocument(
        parser="pdf",
        confidence="high",
        blocks=[
            ParsedBlock(
                kind="paragraph",
                text="Общая площадь объекта 8426,7 кв.м.",
                path="/page/4/p/3",
                page_no=4,
            )
        ],
        warnings=[],
    )
    text = parsed.extraction_text("ТЗ.pdf")
    assert "[DOCUMENT=ТЗ.pdf" in text
    assert "PATH=/page/4/p/3" in text
    assert "page=4" in text
    assert "8426,7" in text


def test_chunk_extraction_text_keeps_source_markers_and_respects_limit():
    text = "[DOCUMENT=7:ТЗ.docx PATH=/p/1 KIND=paragraph]\n" + ("Требование по графику. " * 90)

    chunks = chunk_extraction_text(text, max_chars=180)

    assert len(chunks) > 1
    assert all(len(chunk) <= 180 for chunk in chunks)
    assert all(chunk.startswith("[DOCUMENT=7:ТЗ.docx PATH=/p/1 KIND=paragraph]\n") for chunk in chunks)
    assert "Требование по графику." in " ".join(chunks)


def test_merge_extractions_records_conflicting_values_across_chunks():
    from app.main import _merge_extractions

    first = FieldEvidence(
        key="area_m2", label="Площадь", value="1200", unit="м²",
        source_document="1:dogovor.docx", source_location="/p/1", source_fragment="Площадь 1200 м²",
    )
    second = FieldEvidence(
        key="area_m2", label="Площадь", value="1500", unit="м²",
        source_document="2:tehnicheskoe-zadanie.docx", source_location="/p/2", source_fragment="Площадь 1500 м²",
    )

    result = _merge_extractions([
        DealExtraction(fields=[first], missing_fields=["area_m2", "schedule"]),
        DealExtraction(fields=[second], missing_fields=["schedule"]),
    ])

    assert len(result.fields) == 2
    assert result.missing_fields == ["schedule"]
    assert any(item.get("key") == "area_m2" for item in result.contradictions)


def test_xlsx_parser_obeys_configured_row_limit():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["header"])
    sheet.append(["row 1"])
    sheet.append(["row 2"])
    binary = BytesIO()
    workbook.save(binary)

    from app.document_processing import parse_xlsx

    parsed = parse_xlsx(binary.getvalue(), max_rows=2)
    assert len(parsed.blocks[0].rows) == 2
    assert parsed.warnings


def test_xlsx_table_text_labels_values_by_column_without_pipe_separators():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Показатель", "Значение"])
    sheet.append(["Площадь", "8426,7 м²"])
    binary = BytesIO()
    workbook.save(binary)

    from app.document_processing import parse_xlsx

    parsed = parse_xlsx(binary.getvalue(), max_rows=10)

    assert parsed.blocks[0].text == "Строка 2: Показатель: Площадь; Значение: 8426,7 м²"
    assert " | " not in parsed.blocks[0].text
    assert parsed.blocks[0].rows == [["Показатель", "Значение"], ["Площадь", "8426,7 м²"]]
    assert parsed.warnings and "первая строка принята" in parsed.warnings[0]

    table_requirement = FieldEvidence(
        key="area_m2", label="Площадь", value="8426,7", unit="м²",
        source_location="/sheet/1", source_fragment="Значение: 8426,7 м²",
    )
    assert _field_source_is_valid(table_requirement, {"blocks": [
        {"path": parsed.blocks[0].path, "text": parsed.blocks[0].text}
    ]})


def test_docx_table_text_labels_values_and_keeps_structured_rows():
    from docx import Document

    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Показатель"
    table.cell(0, 1).text = "Значение"
    table.cell(1, 0).text = "График"
    table.cell(1, 1).text = "Ежедневно"
    binary = BytesIO()
    document.save(binary)

    from app.document_processing import parse_docx

    parsed = parse_docx(binary.getvalue())

    assert parsed.blocks[0].text == "Строка 2: Показатель: График; Значение: Ежедневно"
    assert parsed.blocks[0].rows == [["Показатель", "Значение"], ["График", "Ежедневно"]]
    assert parsed.warnings and "первая строка принята" in parsed.warnings[0]


def test_markdown_tables_are_structured_and_text_does_not_use_pipe_delimiters():
    from app.document_processing import parse_document

    parsed = parse_document(
        "requirements.md", "text/markdown",
        "| Параметр | Значение |\n| --- | --- |\n| Площадь | 1200 м² |".encode(), 100,
    )

    table = next(block for block in parsed.blocks if block.kind == "table")
    assert table.rows == [["Параметр", "Значение"], ["Площадь", "1200 м²"]]
    assert "Параметр: Площадь" in table.text
    assert " | " not in table.text


def test_html_parser_ignores_script_and_extracts_tables():
    from app.document_processing import parse_document

    parsed = parse_document(
        "requirements.html", "text/html",
        b"<h1>Requirements</h1><script>alert(1)</script><table><tr><th>Field</th><th>Value</th></tr><tr><td>Area</td><td>1200</td></tr></table>",
        100,
    )

    assert "Requirements" in parsed.plain_text
    assert "alert" not in parsed.plain_text
    table = next(block for block in parsed.blocks if block.kind == "table")
    assert table.rows == [["Field", "Value"], ["Area", "1200"]]


def test_rtf_parser_preserves_cyrillic_paragraphs():
    from app.document_processing import parse_document

    parsed = parse_document("requirements.rtf", "application/rtf", "{\\rtf1\\ansi\\ansicpg1251 Требование\\par}".encode("cp1251"), 100)

    assert parsed.parser == "rtf"
    assert any("Требование" in block.text for block in parsed.blocks)


def test_mht_parser_extracts_html_part():
    from email.message import EmailMessage

    from app.document_processing import parse_document

    message = EmailMessage()
    message.set_type("multipart/related")
    message.add_alternative("<p>Требование из веб-архива</p>", subtype="html", charset="utf-8")
    parsed = parse_document("requirements.mht", "multipart/related", message.as_bytes(), 100)

    assert parsed.parser == "mht"
    assert "Требование из веб-архива" in parsed.plain_text
    assert any("изображения не распознавались" in warning for warning in parsed.warnings)


def test_pptx_parser_keeps_slide_number_and_table_rows():
    from pptx import Presentation
    from pptx.util import Inches

    from app.document_processing import parse_document

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Требование на слайде"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(5), Inches(1)).table
    table.cell(0, 0).text = "Параметр"
    table.cell(0, 1).text = "Значение"
    table.cell(1, 0).text = "Площадь"
    table.cell(1, 1).text = "1200 м²"
    binary = BytesIO()
    presentation.save(binary)

    parsed = parse_document("requirements.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation", binary.getvalue(), 100)

    assert parsed.parser == "pptx"
    assert any(block.page_no == 1 and "Требование на слайде" in block.text for block in parsed.blocks)
    table_block = next(block for block in parsed.blocks if block.kind == "table")
    assert table_block.rows[1] == ["Площадь", "1200 м²"]


def test_pdf_parser_detects_tables_as_structured_blocks():
    import fitz

    from app.document_processing import parse_pdf

    document = fitz.open()
    page = document.new_page()
    for x in (50, 200, 350):
        page.draw_line((x, 50), (x, 140))
    for y in (50, 80, 110, 140):
        page.draw_line((50, y), (350, y))
    for x, y, value in (
        (55, 70, "Field"), (205, 70, "Value"),
        (55, 100, "Area"), (205, 100, "1200 m2"),
    ):
        page.insert_text((x, y), value)
    parsed = parse_pdf(document.tobytes())
    document.close()

    table = next(block for block in parsed.blocks if block.kind == "table")
    assert table.rows[1] == ["Area", "1200 m2"]
    assert "Field: Area" in table.text


def test_field_source_requires_matching_path_and_verbatim_fragment():
    document = {
        "blocks": [{"path": "/page/4/p/3", "page_no": 4, "text": "Общая площадь 8426,7 кв. м."}]
    }
    verified = FieldEvidence(
        key="area_m2",
        label="Площадь",
        value="8426,7",
        source_location="/page/4/p/3 page 4",
        source_fragment="Общая площадь 8426,7 кв. м.",
    )
    fabricated_quote = verified.model_copy(update={"source_fragment": "Общая площадь 9000 кв. м."})
    wrong_location = verified.model_copy(update={"source_location": "/page/3/p/1 page 3"})

    assert _field_source_is_valid(verified, document)
    assert not _field_source_is_valid(fabricated_quote, document)
    assert not _field_source_is_valid(wrong_location, document)


def test_field_evidence_normalizes_numeric_model_values():
    field = FieldEvidence(key="area_m2", label="Площадь", value=1200, unit="м²")

    assert field.value == "1200"


def test_field_evidence_normalizes_qualitative_confidence_from_mimo():
    high = FieldEvidence(key="area_m2", label="Площадь", value="1200", confidence="высокая", status=None)
    medium = FieldEvidence(key="schedule", label="График", value="ежедневно", confidence="средняя")
    low = FieldEvidence(key="staff", label="Персонал", value="4", confidence="низкая")
    unknown = FieldEvidence(key="unknown", label="Неизвестно", confidence="не уверен")

    assert high.confidence == 0.95
    assert high.status == "extracted"
    assert medium.confidence == 0.7
    assert low.confidence == 0.4
    assert unknown.confidence is None


def test_deal_extraction_normalizes_non_list_model_contradictions():
    empty = DealExtraction.model_validate({"contradictions": ": "})
    text = DealExtraction.model_validate({"contradictions": "Площадь различается в двух документах"})
    items = DealExtraction.model_validate({"contradictions": ["Разные сроки оплаты"]})

    assert empty.contradictions == []
    assert text.contradictions == [{"description": "Площадь различается в двух документах"}]
    assert items.contradictions == [{"description": "Разные сроки оплаты"}]
