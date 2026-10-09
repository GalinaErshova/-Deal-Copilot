from io import BytesIO

from docx import Document as WordDocument
from fastapi.testclient import TestClient

from app.config import settings
from app.main import app

client = TestClient(app)


def create_deal() -> int:
    response = client.post("/api/deals", params={"title": "API validation test"})
    assert response.status_code == 200
    return response.json()["id"]


def test_settings_endpoint_exposes_public_defaults_only():
    response = client.get("/api/settings")
    assert response.status_code == 200
    payload = response.json()
    assert payload["calculation_defaults"]["area_m2"] == settings.default_area_m2
    assert payload["accepted_upload_extensions"] == settings.parsed_upload_extensions
    assert "mimo_api_key" not in payload


def test_calculation_requires_confirmed_area_and_matching_value():
    deal_id = create_deal()
    defaults = settings.calculation_defaults

    assert client.post(f"/api/deals/{deal_id}/calculate", json=defaults).status_code == 409

    response = client.post(f"/api/deals/{deal_id}/fields/area", json={"value": defaults["area_m2"]})
    assert response.status_code == 200
    field_id = response.json()["id"]
    assert client.patch(f"/api/fields/{field_id}", json={"value": defaults["area_m2"], "confirmed": True}).status_code == 200

    changed_area = defaults | {"area_m2": defaults["area_m2"] + 1}
    assert client.post(f"/api/deals/{deal_id}/calculate", json=changed_area).status_code == 409

    response = client.post(f"/api/deals/{deal_id}/calculate", json=defaults)
    assert response.status_code == 200
    assert response.json()["decision"] in {"BID", "BID WITH CONDITIONS", "NO BID"}


def test_calculation_rejects_invalid_numeric_input():
    deal_id = create_deal()
    defaults = settings.calculation_defaults
    response = client.post(f"/api/deals/{deal_id}/fields/area", json={"value": defaults["area_m2"]})
    field_id = response.json()["id"]
    client.patch(f"/api/fields/{field_id}", json={"value": defaults["area_m2"], "confirmed": True})

    invalid = defaults | {"productivity_m2_per_shift": 0}
    assert client.post(f"/api/deals/{deal_id}/calculate", json=invalid).status_code == 422


def test_saved_calculation_becomes_history_after_confirmed_area_changes():
    deal_id=create_deal()
    defaults=settings.calculation_defaults
    created=client.post(f"/api/deals/{deal_id}/fields/area",json={"value":defaults["area_m2"]})
    field_id=created.json()["id"]
    assert client.patch(f"/api/fields/{field_id}",json={"confirmed":True}).status_code==200
    calculated=client.post(f"/api/deals/{deal_id}/calculate",json=defaults)
    assert calculated.status_code==200
    assert client.get(f"/api/deals/{deal_id}/calculations").json()[0]["is_current"] is True
    assert client.patch(f"/api/fields/{field_id}",json={"value":str(defaults["area_m2"]+1),
                                                         "confirmed":True}).status_code==200
    history=client.get(f"/api/deals/{deal_id}/calculations").json()
    assert history[0]["id"]==calculated.json()["calculation_id"]
    assert history[0]["is_current"] is False
    assert client.post(f"/api/deals/{deal_id}/proposal",json={}).status_code==409


def test_invalid_formula_is_rejected_without_replacing_saved_formula():
    original=client.get("/api/formulas").json()
    submitted=[{"key":item["key"],"expression":item["expression"]} for item in original]
    next(item for item in submitted if item["key"]=="labor_hours_month")["expression"]="1e309"
    response=client.put("/api/formulas",json={"formulas":submitted})
    assert response.status_code==422
    assert client.get("/api/formulas").json()==original


def test_document_extraction_over_limit_is_explicitly_rejected(monkeypatch):
    from app.main import settings as app_settings

    deal_id=create_deal()
    document=WordDocument()
    document.add_paragraph("Условия уборки "*30)
    binary=BytesIO()
    document.save(binary)
    assert client.post(f"/api/deals/{deal_id}/documents",files=[
        ("files",("oversize.docx",binary.getvalue(),
                  "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
    ]).status_code==200
    monkeypatch.setattr(app_settings,"max_extraction_chars",50)
    response=client.post(f"/api/deals/{deal_id}/process")
    assert response.status_code==413
    assert "превышает лимит" in response.json()["detail"]
    pipeline=client.get(f"/api/deals/{deal_id}/pipeline").json()
    assert pipeline[0]["status"]=="failed"
    assert pipeline[0]["steps"][-1]["output"]["omitted_characters"]>0


def test_breakdown_totals_term_and_manual_line_in_proposal():
    from app.db import SessionLocal
    from app.models import ExtractedField

    deal_id=create_deal()
    with SessionLocal() as db:
        db.add(ExtractedField(deal_id=deal_id,key="contract_months",label="Срок",
                              value="10",unit="мес.",confirmed=True,status="user_confirmed"))
        db.commit()
    created=client.post(f"/api/deals/{deal_id}/area-components/manual",json={
        "address":"Объект А","work_type":"Уборка","area_m2":100,
        "schedule_mode":"daily","productivity_m2_per_shift":500,
        "price_per_m2_month":180})
    assert created.status_code==200
    component=client.get(f"/api/deals/{deal_id}/area-components").json()[0]
    assumptions=settings.calculation_defaults|{"contract_months":10}
    body={"confirmed":True,"assumptions":assumptions,"components":[{
        "id":component["id"],"area_m2":100,"productivity_m2_per_shift":500,
        "schedule_mode":"daily","price_per_m2_month":180}]}
    assert client.post(f"/api/deals/{deal_id}/calculate-breakdown",json={
        **body,"assumptions":assumptions|{"contract_months":12}}).status_code==409
    calculated=client.post(f"/api/deals/{deal_id}/calculate-breakdown",json=body)
    assert calculated.status_code==200,calculated.text
    result=calculated.json()
    assert result["components"][0]["monthly_price"]==result["revenue_with_vat"]
    assert result["components"][0]["contract_price"]==180000
    assert result["components"][0]["schedule_label"]=="Каждый рабочий день"
    proposal=client.post(f"/api/deals/{deal_id}/proposal",json={})
    assert proposal.status_code==200,proposal.text
    document=WordDocument(BytesIO(proposal.content))
    all_text="\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Срок оказания услуг: 10 мес." in all_text
    assert "180 000,00" in all_text
    assert client.delete(f"/api/deals/{deal_id}/area-components/manual/{created.json()['id']}").status_code==200
    assert client.get(f"/api/deals/{deal_id}/calculations").json()[0]["is_current"] is False
    assert client.post(f"/api/deals/{deal_id}/proposal",json={}).status_code==409


def test_document_catalog_price_and_changed_schedule_use_current_values():
    import json

    from app.db import SessionLocal
    from app.models import Document

    deal_id=create_deal()
    blocks=[{"kind":"table","path":"/table/1","text":
             "Комплексная уборка помещений. Ежедневно: Х",
             "rows":[["Объект","Площадь","Ежедневно"],["Офис тест","100","Х"]]}]
    with SessionLocal() as db:
        db.add(Document(deal_id=deal_id,filename="scope.docx",content_type="application/docx",
                        file_path="unused",parse_status="parsed",parser="docx",
                        extracted_text="Уборка",parse_json=json.dumps({"blocks":blocks},ensure_ascii=False)))
        db.commit()
    catalog=client.post("/api/price-list",json={"name":"Тестовый тариф","area_type":"Помещения",
        "work_type":"Комплексная уборка помещений","price_per_m2_month":180})
    assert catalog.status_code==200
    catalog_id=catalog.json()["id"]
    source=client.get(f"/api/deals/{deal_id}/area-components").json()[0]
    assert source["curation_status"]=="verified"
    assert source["schedule_mode"]=="daily"
    assert source["price_list_item_id"]==catalog_id
    assert client.patch(f"/api/price-list/{catalog_id}",json={"price_per_m2_month":200}).status_code==200
    body={"confirmed":True,"assumptions":settings.calculation_defaults,"components":[{
        "id":source["id"],"area_m2":100,"productivity_m2_per_shift":500,
        "schedule_mode":"monthly","price_per_m2_month":180}]}
    response=client.post(f"/api/deals/{deal_id}/calculate-breakdown",json=body)
    assert response.status_code==200,response.text
    line=response.json()["components"][0]
    assert line["price_per_m2_month"]==200
    assert line["monthly_price"]==response.json()["revenue_with_vat"]==20000
    assert line["schedule_label"]=="Ежемесячно"
    assert line["original_schedule_label"].startswith("Основные операции: Ежедневно")
    assert line["schedule_manually_changed"] is True
    proposal=client.post(f"/api/deals/{deal_id}/proposal",json={})
    assert proposal.status_code==200,proposal.text
    document=WordDocument(BytesIO(proposal.content))
    assert "Ежемесячно" in " ".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)


def test_edited_line_formulas_change_aggregate_labor_and_revenue():
    from app.db import SessionLocal
    from app.models import CalculationFormula

    deal_id=create_deal()
    created=client.post(f"/api/deals/{deal_id}/area-components/manual",json={
        "address":"Объект формул","work_type":"Уборка","area_m2":100,
        "schedule_mode":"daily","productivity_m2_per_shift":500,
        "price_per_m2_month":180})
    assert created.status_code==200
    body={"confirmed":True,"assumptions":settings.calculation_defaults,"components":[{
        "id":created.json()["component_id"],"area_m2":100,
        "productivity_m2_per_shift":500,"schedule_mode":"daily",
        "price_per_m2_month":180}]}
    baseline=client.post(f"/api/deals/{deal_id}/calculate-breakdown",json=body)
    assert baseline.status_code==200,baseline.text
    with SessionLocal() as db:
        labor=db.get(CalculationFormula,"line_labor_hours_month")
        price=db.get(CalculationFormula,"line_monthly_price")
        old_labor,old_price=labor.expression,price.expression
        labor.expression=f"({old_labor}) * 2"
        price.expression=f"({old_price}) * 0.9"
        db.commit()
    try:
        changed=client.post(f"/api/deals/{deal_id}/calculate-breakdown",json=body)
        assert changed.status_code==200,changed.text
        original,updated=baseline.json(),changed.json()
        assert updated["labor_hours_month"]==updated["components"][0]["labor_hours_month"]
        assert updated["labor_hours_month"]==original["labor_hours_month"]*2
        assert updated["revenue_with_vat"]==updated["components"][0]["monthly_price"]
        assert updated["revenue_with_vat"]==original["revenue_with_vat"]*0.9
        assert updated["direct_cost"]>original["direct_cost"]
        assert updated["margin"]<original["margin"]
    finally:
        with SessionLocal() as db:
            db.get(CalculationFormula,"line_labor_hours_month").expression=old_labor
            db.get(CalculationFormula,"line_monthly_price").expression=old_price
            db.commit()


def test_upload_rejects_unsupported_extensions():
    deal_id = create_deal()
    response = client.post(
        f"/api/deals/{deal_id}/documents",
        files=[("files", ("payload.exe", b"not an accepted document", "application/octet-stream"))],
    )
    assert response.status_code == 415


def test_docx_upload_processing_completes_in_mock_mode_without_invented_fields(monkeypatch):
    deal_id = create_deal()
    monkeypatch.setattr(
        "app.main.gateway.structured",
        lambda **kwargs: kwargs["schema"].model_validate(
            {"fields": [], "missing_fields": [], "contradictions": []}
        ),
    )
    document = WordDocument()
    document.add_paragraph("Техническое задание на уборку офисного здания.")
    binary = BytesIO()
    document.save(binary)

    uploaded = client.post(
        f"/api/deals/{deal_id}/documents",
        files=[("files", ("requirements.docx", binary.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
    )
    assert uploaded.status_code == 200

    processed = client.post(f"/api/deals/{deal_id}/process")
    assert processed.status_code == 200
    assert processed.json()["fields"] == []
    assert client.get(f"/api/deals/{deal_id}/fields").json() == []

    pipeline = client.get(f"/api/deals/{deal_id}/pipeline").json()
    assert all(run["status"] == "success" for run in pipeline)


def test_uploaded_documents_remain_available_when_model_extraction_fails(monkeypatch):
    deal_id = create_deal()
    document = WordDocument()
    document.add_paragraph("Техническое задание на уборку офисного здания.")
    binary = BytesIO()
    document.save(binary)
    uploaded = client.post(
        f"/api/deals/{deal_id}/documents",
        files=[("files", ("requirements.docx", binary.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
    )
    assert uploaded.status_code == 200

    def fail_extraction(**_kwargs):
        raise RuntimeError("model unavailable")

    monkeypatch.setattr("app.main.gateway.structured", fail_extraction)
    processed = client.post(f"/api/deals/{deal_id}/process")

    assert processed.status_code == 502
    documents = client.get(f"/api/deals/{deal_id}/documents").json()
    assert len(documents) == 1
    assert documents[0]["status"] == "parsed"
    original = client.get(f"/api/documents/{documents[0]['id']}/original")
    assert original.status_code == 200
    parsed = client.get(f"/api/documents/{documents[0]['id']}/parsed").json()
    assert "Техническое задание" in parsed["blocks"][0]["text"]
    deals = client.get("/api/deals").json()
    assert next(deal for deal in deals if deal["id"] == deal_id)["document_count"] == 1


def test_model_qualitative_confidence_does_not_fail_document_pipeline(monkeypatch):
    deal_id = create_deal()
    document = WordDocument()
    document.add_paragraph("Техническое задание на уборку офисного здания.")
    binary = BytesIO()
    document.save(binary)
    uploaded = client.post(
        f"/api/deals/{deal_id}/documents",
        files=[("files", ("requirements.docx", binary.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
    )
    assert uploaded.status_code == 200

    model_reply = {
        "fields": [{"key": "area_m2", "label": "Площадь", "value": "1200", "unit": "м²", "confidence": "высокая", "status": None}],
        "missing_fields": [],
        "contradictions": ": ",
    }
    monkeypatch.setattr(
        "app.main.gateway.structured",
        lambda **kwargs: kwargs["schema"].model_validate(model_reply),
    )

    processed = client.post(f"/api/deals/{deal_id}/process")

    assert processed.status_code == 200
    fields = client.get(f"/api/deals/{deal_id}/fields").json()
    assert fields[0]["confidence"] == 0.95
    pipeline = client.get(f"/api/deals/{deal_id}/pipeline").json()
    assert pipeline[0]["status"] == "success"
    extraction_step = next(step for step in pipeline[0]["steps"] if step["name"] == "ai_extraction")
    assert extraction_step["output"]["contradictions"] == []
